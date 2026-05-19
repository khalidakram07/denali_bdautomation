"""
services/template_engine.py - .docx email template loader and renderer.

Templates live in templates/*.docx. Each template can contain:
  - {placeholders} - direct substitution from contact + opportunity data
  - [BRACKETED INSTRUCTIONS] - filled in by the AI using the Full Text article

Workflow:
  1. list_templates()        - discover templates for the UI dropdown
  2. load_template_text()    - read a .docx, return plain text
  3. substitute_placeholders() - swap {var} tokens
  4. The AI step in ai_engine.py uses the substituted template + Full Text
     to fill in the [INSTRUCTIONS] sections.
"""

import logging
import re
from pathlib import Path
from typing import Optional

from docx import Document

log = logging.getLogger(__name__)

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"


def _humanize(stem: str) -> str:
    """first_touch_dr_chen -> 'First Touch Dr Chen'"""
    return stem.replace("_", " ").replace("-", " ").title()


# Drive-sourced filenames are prefixed so they don't collide with local files
# AND so the engine knows where to load them from.
DRIVE_PREFIX = "drive:"


def list_templates(force_refresh: bool = False) -> list[dict]:
    """
    Merged list of templates from:
      1. Google Drive folder (if GOOGLE_DRIVE_TEMPLATES_FOLDER_ID is set)
      2. Local templates/ folder

    Drive entries are prefixed with 'drive:' in the filename, and labelled
    with a 🔗 in the display name so it's obvious in the UI dropdown.
    """
    out = []

    # 1. Drive templates (if Drive is configured)
    try:
        from services.drive_templates import list_drive_templates
        for t in list_drive_templates(force_refresh=force_refresh):
            out.append({
                "filename":     DRIVE_PREFIX + t["filename"],
                "display_name": "🔗 " + t["display_name"],
                "size_kb":      None,
                "source":       "drive",
                "modified":     t.get("modified_time"),
            })
    except Exception as e:
        log.warning("Drive templates unavailable: %s", e)

    # 2. Local templates
    if TEMPLATES_DIR.exists():
        for path in sorted(TEMPLATES_DIR.glob("*.docx")):
            if path.name.startswith("~$"):
                continue
            out.append({
                "filename":     path.name,
                "display_name": _humanize(path.stem),
                "size_kb":      round(path.stat().st_size / 1024, 1),
                "source":       "local",
                "modified":     None,
            })
    return out


def load_template_text(filename: str) -> Optional[str]:
    """
    Read a template's text. Resolves Drive- or local-sourced filenames.
    Returns None if missing.
    """
    # Drive-sourced template
    if filename.startswith(DRIVE_PREFIX):
        from services.drive_templates import load_drive_template_text
        return load_drive_template_text(filename[len(DRIVE_PREFIX):])

    # Local file
    if not filename.lower().endswith(".docx"):
        return None
    path = TEMPLATES_DIR / filename
    if not path.exists() or not path.is_file():
        return None
    try:
        doc = Document(str(path))
    except Exception as e:
        log.error("Failed to open template %s: %s", filename, e)
        return None
    return "\n".join(p.text for p in doc.paragraphs)


def refresh_templates_cache() -> int:
    """Clear the Drive template cache. Returns # of items cleared."""
    try:
        from services.drive_templates import refresh_cache
        return refresh_cache()
    except Exception:
        return 0


PLACEHOLDER_RE = re.compile(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}")


def substitute_placeholders(template_text: str, vars: dict) -> str:
    """
    Replace {placeholder} tokens with values from `vars`.
    Unknown placeholders are left untouched (so the AI can flag them).
    """
    def repl(match):
        key = match.group(1)
        if key in vars and vars[key] is not None and vars[key] != "":
            return str(vars[key])
        return match.group(0)
    return PLACEHOLDER_RE.sub(repl, template_text)


def has_ai_instructions(template_text: str) -> bool:
    """True if the template contains [BRACKETED] instructions for the AI to fill."""
    return bool(re.search(r"\[[A-Z][^\]]{3,}\]", template_text))
